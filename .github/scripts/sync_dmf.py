#!/usr/bin/env python3
"""
sync_dmf.py - Headless ETL Pipeline for DMF Hydrocarbon Data Sync
Executed via GitHub Actions (scheduled / workflow_dispatch) or CLI.
Extracts, validates, cleans, maps, and normalizes domestic production & sales data.
Outputs clean static JSON for Next.js frontend and Excel flat tables for PTIT analysts.
"""

import os
import sys
import re
import io
import json
import glob
import time
import hashlib
import datetime
import math
from pathlib import Path

import openpyxl
import pandas as pd

# Ensure UTF-8 stdout for Windows consoles
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def clean_for_json(data):
    """Recursively clean data structures to ensure 100% valid JSON (no NaN or Infinity)."""
    if isinstance(data, dict):
        return {k: clean_for_json(v) for k, v in data.items()}
    elif isinstance(data, list):
        return [clean_for_json(v) for v in data]
    elif isinstance(data, float):
        if math.isnan(data) or math.isinf(data):
            return 0.0
        return data
    elif pd.isna(data):
        return None
    return data

# Path configurations
SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent.parent

# Input Master Files
MASTER_MAPPING_FILE = REPO_ROOT / "master_mapping.xlsx"
SALE_MAPPING_FILE = REPO_ROOT / "sale_master_mapping.xlsx"
FANG_MASTER_JSON = REPO_ROOT / "fang_production_master.json"
FANG_MASTER_XLSX = REPO_ROOT / "fang_production_master.xlsx"

# Output Locations
DATA_DIR = REPO_ROOT / "data"
FRONTEND_DATA_DIR = REPO_ROOT / "frontend" / "public" / "data"
EXCEL_OUTPUT_DIR = REPO_ROOT / "output"

PRODUCTION_JSON = DATA_DIR / "production_master.json"
SALES_JSON = DATA_DIR / "sales_master.json"
MANIFEST_JSON = DATA_DIR / "sync_manifest.json"

EXCEL_PROD_OUTPUT = EXCEL_OUTPUT_DIR / "petroleum_production_flat_table.xlsx"
EXCEL_SALE_OUTPUT = EXCEL_OUTPUT_DIR / "petroleum_sale_flat_table.xlsx"

# Import live stream module if in path
sys.path.insert(0, str(REPO_ROOT))
try:
    import dmf_live_stream as dmf
except ImportError:
    dmf = None

# Thai Month References
THAI_MONTHS = [
    'มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
    'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม'
]
MONTH_ORDER = {m: i + 1 for i, m in enumerate(THAI_MONTHS)}
DAYS_IN_MONTH = {
    'มกราคม': 31, 'กุมภาพันธ์': 28, 'มีนาคม': 31, 'เมษายน': 30,
    'พฤษภาคม': 31, 'มิถุนายน': 30, 'กรกฎาคม': 31, 'สิงหาคม': 31,
    'กันยายน': 30, 'ตุลาคม': 31, 'พฤศจิกายน': 30, 'ธันวาคม': 31
}

def get_days_in_month(month_name, year_val=2569):
    """Accurate days count with leap-year (29 Feb) handling."""
    try:
        y_int = int(year_val)
        y_ce = y_int - 543 if y_int > 2400 else y_int
    except Exception:
        y_ce = 2026
    m_clean = str(month_name).strip()
    if m_clean == 'กุมภาพันธ์':
        if (y_ce % 4 == 0 and y_ce % 100 != 0) or (y_ce % 400 == 0):
            return 29
        return 28
    return DAYS_IN_MONTH.get(m_clean, 30)

def clean_number(val):
    if val is None or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        clean_str = re.sub(r'[, \t\n]', '', val)
        try:
            return float(clean_str)
        except ValueError:
            return 0.0
    return 0.0

def normalize_text(text):
    if not text:
        return ''
    return re.sub(r'\s+', ' ', str(text)).strip()

def get_merged_cell_value(ws, row, col):
    for rng in ws.merged_cells.ranges:
        if row >= rng.min_row and row <= rng.max_row and col >= rng.min_col and col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value

def load_master_mapping():
    if os.path.exists(MASTER_MAPPING_FILE):
        return pd.read_excel(MASTER_MAPPING_FILE)
    cols = [
        "Lookup_Key", "พื้นที่", "แปลง_ไฟล์ดิบ", "แหล่ง_ไฟล์ดิบ",
        "PTIT_Region", "PTIT_Operator_Field", "PTIT_Order",
        "ผู้ดำเนินการ", "แอ่งปิโตรเลียม", "ประเภทสัญญา"
    ]
    return pd.DataFrame(columns=cols)

def load_sale_master_mapping():
    if os.path.exists(SALE_MAPPING_FILE):
        return pd.read_excel(SALE_MAPPING_FILE)
    cols = ["แหล่ง_ไฟล์ดิบ", "พื้นที่", "ผู้ดำเนินการ", "แอ่งปิโตรเลียม", "ประเภทสัญญา", "หมายเหตุ"]
    return pd.DataFrame(columns=cols)

def load_fang_master():
    default_data = {
        "2569": {
            "มกราคม": 0.0, "กุมภาพันธ์": 0.0, "มีนาคม": 0.0, "เมษายน": 0.0,
            "พฤษภาคม": 0.0, "มิถุนายน": 610.4, "กรกฎาคม": 0.0, "สิงหาคม": 0.0,
            "กันยายน": 0.0, "ตุลาคม": 0.0, "พฤศจิกายน": 0.0, "ธันวาคม": 0.0
        }
    }
    if os.path.exists(FANG_MASTER_XLSX):
        try:
            wb = openpyxl.load_workbook(FANG_MASTER_XLSX, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) > 1:
                res = {}
                hdr = [str(c or '').strip() for c in rows[0]]
                idx_y = hdr.index("ปี") if "ปี" in hdr else 0
                idx_m = hdr.index("เดือน") if "เดือน" in hdr else 2
                idx_val = len(hdr) - 1
                for i, h in enumerate(hdr):
                    if "ค่าน้ำมันดิบ" in h or "BPD" in h:
                        idx_val = i
                        break
                for r in rows[1:]:
                    if r[idx_y] is not None and r[idx_m] is not None:
                        y = str(r[idx_y]).strip()
                        m = str(r[idx_m]).strip()
                        v = float(r[idx_val]) if r[idx_val] is not None and str(r[idx_val]).replace('.', '', 1).isdigit() else 0.0
                        if y not in res:
                            res[y] = {}
                        res[y][m] = v
                if res:
                    return res
        except Exception:
            pass

    if os.path.exists(FANG_MASTER_JSON):
        try:
            with open(FANG_MASTER_JSON, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return default_data

def inject_fang_to_dataframe(df_source):
    if df_source is None or df_source.empty:
        return df_source
    df_clean = df_source[df_source.get('ผู้ดำเนินการ', '') != 'Defence Energy Department'].copy()
    fang_data = load_fang_master()
    fang_rows = []
    years_in_data = [str(y) for y in df_clean['ปี'].dropna().unique().tolist()] if 'ปี' in df_clean.columns else ['2569']

    for y_str in years_in_data:
        y_fang = fang_data.get(y_str, {})
        for m_name, m_ord in MONTH_ORDER.items():
            val = float(y_fang.get(m_name, 0.0))
            if val > 0:
                y_val = int(y_str) if y_str.isdigit() else y_str
                days_cnt = get_days_in_month(m_name, y_val)
                fang_rows.append({
                    'ปี': y_val,
                    'เดือน': m_name,
                    'ลำดับเดือน': m_ord,
                    'พื้นที่': 'บนบก',
                    'PTIT_Region': 'Onshore',
                    'PTIT_Operator_Field': 'Defence Energy Department / Fang',
                    'PTIT_Order': 13,
                    'ผู้ดำเนินการ': 'Defence Energy Department',
                    'แอ่งปิโตรเลียม': 'Fang Basin',
                    'ประเภทสัญญา': 'DEDP',
                    'ก๊าซธรรมชาติ (ล้านลบ.ฟุต/วัน)': 0.0,
                    'ก๊าซธรรมชาติ_เทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)': 0.0,
                    'ก๊าซธรรมชาติเหลว (บาร์เรล/วัน)': 0.0,
                    'ก๊าซธรรมชาติเหลว_เทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)': 0.0,
                    'น้ำมันดิบ (บาร์เรล/วัน)': val,
                    'รวมเทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)': val,
                    'จำนวนวันที่ผลิต': days_cnt,
                    'หมายเหตุ': 'ข้อมูลจากกรมการพลังงานทหาร (DEDP)',
                    'แปลง_ไฟล์ดิบ': 'Fang',
                    'แหล่ง_ไฟล์ดิบ': 'Fang',
                    'Lookup_Key': 'บนบก_Fang_Fang',
                    'ไฟล์ที่มา': 'DEDP_Fang_Verified'
                })

    if fang_rows:
        df_fang = pd.DataFrame(fang_rows)
        for col in df_clean.columns:
            if col not in df_fang.columns:
                df_fang[col] = None
        return pd.concat([df_clean, df_fang], ignore_index=True)
    return df_clean

def parse_excel_file(file_input, filename_label):
    wb = openpyxl.load_workbook(file_input, data_only=True)
    ws = wb.active

    header_row_idx = None
    for r in range(1, min(30, ws.max_row + 1)):
        row_values = [normalize_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any('พื้นที่' in v for v in row_values) and any('แปลง' in v for v in row_values):
            header_row_idx = r
            break

    if header_row_idx is None:
        raise ValueError(f"ไม่พบหัวตารางในไฟล์ {filename_label}")

    month_name, year_val = '', '2569'
    for r in range(1, header_row_idx):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value or '')
            match = re.search(r'เดือน\s*([^\s]+)\s*ปี\s*(\d+)', val)
            if match:
                month_name, year_val = match.group(1), match.group(2)
                break
        if month_name:
            break

    col_map = {}
    for c in range(1, ws.max_column + 1):
        top = get_merged_cell_value(ws, header_row_idx - 1, c)
        top = normalize_text(top)
        sub = normalize_text(ws.cell(header_row_idx, c).value)
        full = f"{top} {sub}".strip()

        if 'พื้นที่' in sub:
            col_map['area'] = c
        elif 'แปลง' in sub:
            col_map['block'] = c
        elif 'แหล่ง' in sub:
            col_map['field'] = c
        elif 'จำนวนวัน' in sub or 'วันผลิต' in sub or 'days' in sub.lower():
            col_map['days'] = c
        elif 'หมายเหตุ' in sub:
            col_map['notes'] = c
        elif 'เทียบเท่า' in full or 'เทียบเท่า' in sub:
            if 'เหลว' in full or 'คอนเดนเสท' in full:
                col_map['cond_boed'] = c
            elif 'ก๊าซ' in full:
                col_map['gas_boed'] = c
        elif 'น้ำมันดิบ' in full and 'เทียบเท่า' not in full:
            col_map['crude_bpd'] = c
        elif 'ก๊าซ' in full and any(u in full for u in ['ล้านลบ', 'mmscfd']):
            col_map['gas_mmscfd'] = c
        elif ('เหลว' in full or 'คอนเดนเสท' in full) and 'เทียบเท่า' not in full:
            col_map['cond_bpd'] = c

    area_col = col_map.get('area', 1)
    block_col = col_map.get('block', 2)
    field_col = col_map.get('field', 3)

    records = []
    for r in range(header_row_idx + 1, ws.max_row + 1):
        area_val = normalize_text(ws.cell(r, area_col).value)
        if not area_val or any(k in area_val for k in ['รวมอัตรา', 'รายงานปริมาณ', 'หมายเหตุ']):
            continue

        raw_block = normalize_text(ws.cell(r, block_col).value)
        raw_field = normalize_text(ws.cell(r, field_col).value)

        gas_mmscfd = clean_number(ws.cell(r, col_map.get('gas_mmscfd')).value) if 'gas_mmscfd' in col_map else 0.0
        gas_boed = clean_number(ws.cell(r, col_map.get('gas_boed')).value) if 'gas_boed' in col_map else 0.0
        cond_bpd = clean_number(ws.cell(r, col_map.get('cond_bpd')).value) if 'cond_bpd' in col_map else 0.0
        cond_boed = clean_number(ws.cell(r, col_map.get('cond_boed')).value) if 'cond_boed' in col_map else 0.0
        crude_bpd = clean_number(ws.cell(r, col_map.get('crude_bpd')).value) if 'crude_bpd' in col_map else 0.0
        total_boed = gas_boed + cond_boed + crude_bpd

        days = clean_number(ws.cell(r, col_map.get('days')).value) if 'days' in col_map else 0.0
        notes = normalize_text(ws.cell(r, col_map.get('notes')).value) if 'notes' in col_map else ''

        records.append({
            'Lookup_Key': f"{area_val}_{raw_block}_{raw_field}",
            'ปี': int(year_val) if year_val.isdigit() else year_val,
            'เดือน': month_name,
            'ลำดับเดือน': MONTH_ORDER.get(month_name, 99),
            'พื้นที่': area_val,
            'แปลง_ไฟล์ดิบ': raw_block,
            'แหล่ง_ไฟล์ดิบ': raw_field,
            'ก๊าซธรรมชาติ (ล้านลบ.ฟุต/วัน)': gas_mmscfd,
            'ก๊าซธรรมชาติ_เทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)': gas_boed,
            'ก๊าซธรรมชาติเหลว (บาร์เรล/วัน)': cond_bpd,
            'ก๊าซธรรมชาติเหลว_เทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)': cond_boed,
            'น้ำมันดิบ (บาร์เรล/วัน)': crude_bpd,
            'รวมเทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)': total_boed,
            'จำนวนวันที่ผลิต': int(days) if days else 0,
            'หมายเหตุ': notes,
            'ไฟล์ที่มา': filename_label
        })

    return pd.DataFrame(records)

def parse_sales_file(file_input, filename_label):
    wb = openpyxl.load_workbook(file_input, data_only=True)
    ws = wb.active

    row2_val = str(ws.cell(2, 1).value or '')
    month_name, year_val = '', '2569'
    for m in THAI_MONTHS:
        if m in row2_val:
            month_name = m
            break
    ymatch = re.search(r'(\d{4})', row2_val)
    if ymatch:
        year_val = ymatch.group(1)

    days_cnt = get_days_in_month(month_name, year_val)
    current_product = None
    records = []

    for r in range(4, ws.max_row + 1):
        c1 = str(ws.cell(r, 1).value or '').strip()
        c2 = ws.cell(r, 2).value
        c3 = ws.cell(r, 3).value
        c4 = ws.cell(r, 4).value
        c5 = ws.cell(r, 5).value

        if c1 in ['ก๊าซธรรมชาติ', 'ก๊าซธรรมชาติเหลว', 'น้ำมันดิบ']:
            current_product = c1
            continue

        if not c1 or c1 in ['แหล่ง', 'รวมทั้งหมด'] or 'ปริมาณ' in str(c2) or 'ล้าน' in str(c2) or 'บาร์เรล' in str(c2):
            continue

        prod_type = current_product
        if 'LPG' in c1:
            prod_type = 'ก๊าซปิโตรเลียมเหลว (LPG)'

        val = clean_number(c4)
        royalty = clean_number(c5)

        mmscf = None
        mmbtu = None
        bbl = None
        kg = None
        rate_daily_vol = None
        rate_daily_heat = None
        rate_daily_bpd = None
        wellhead_price = 0.0
        wellhead_unit = ''
        wellhead_gas_mmscf = None
        heating_val = None

        if prod_type == 'ก๊าซธรรมชาติ':
            mmscf = clean_number(c2)
            mmbtu = clean_number(c3)
            rate_daily_vol = mmscf / days_cnt if days_cnt > 0 else 0.0
            rate_daily_heat = mmbtu / days_cnt if days_cnt > 0 else 0.0
            wellhead_price = val / mmbtu if mmbtu > 0 else 0.0
            wellhead_unit = 'บาท/MMBTU'
            wellhead_gas_mmscf = val / mmscf if mmscf > 0 else 0.0
            heating_val = mmbtu / mmscf if mmscf > 0 else 0.0
            unit_str = 'ล้าน ลบ.ฟุต & ล้านบีทียู'
            v_main = mmscf

        elif prod_type == 'ก๊าซปิโตรเลียมเหลว (LPG)':
            c3_val = clean_number(c3)
            c2_val = clean_number(c2)
            kg = c3_val if c3_val > 0 else c2_val
            rate_daily_vol = kg / days_cnt if days_cnt > 0 else 0.0
            wellhead_price = val / kg if kg > 0 else 0.0
            wellhead_unit = 'บาท/กก.'
            unit_str = 'กิโลกรัม'
            v_main = kg

        elif prod_type in ['ก๊าซธรรมชาติเหลว', 'น้ำมันดิบ']:
            bbl = clean_number(c2)
            rate_daily_bpd = bbl / days_cnt if days_cnt > 0 else 0.0
            wellhead_price = val / bbl if bbl > 0 else 0.0
            wellhead_unit = 'บาท/บาร์เรล'
            unit_str = 'บาร์เรล'
            v_main = bbl

        royalty_pct = (royalty / val * 100) if val > 0 else 0.0

        rec = {
            'ปี': year_val,
            'เดือน': month_name,
            'ลำดับเดือน': MONTH_ORDER.get(month_name, 99),
            'ประเภทปิโตรเลียม': prod_type,
            'แหล่ง_ไฟล์ดิบ': c1,
            'หน่วยปริมาณ': unit_str,
            'ปริมาณการขาย_หน่วยหลัก': v_main,
            'ปริมาณการขาย_MMSCF': mmscf,
            'ปริมาณการขาย_MMBTU': mmbtu,
            'ปริมาณการขาย_บาร์เรล': bbl,
            'ปริมาณการขาย_กิโลกรัม': kg,
            'ปริมาณการขายเฉลี่ย_MMSCFD': rate_daily_vol if prod_type == 'ก๊าซธรรมชาติ' else None,
            'ปริมาณความร้อนเฉลี่ย_MMBTUD': rate_daily_heat if prod_type == 'ก๊าซธรรมชาติ' else None,
            'ปริมาณการขายเฉลี่ย_BPD': rate_daily_bpd if prod_type in ['ก๊าซธรรมชาติเหลว', 'น้ำมันดิบ'] else None,
            'ค่าความร้อน_Heating_Value_BTU_per_SCF': heating_val,
            'มูลค่าการขาย_บาท': val,
            'ค่าภาคหลวง_บาท': royalty,
            'ราคาปากหลุม_Wellhead_Price': wellhead_price,
            'หน่วยราคาปากหลุม': wellhead_unit,
            'ราคาปากหลุม_ก๊าซ_บาทต่อMMSCF': wellhead_gas_mmscf,
            'อัตราค่าภาคหลวงที่แท้จริง_Pct': royalty_pct,
            'ราคาเฉลี่ยต่อหน่วย_บาท': wellhead_price,
            'ไฟล์ที่มา': filename_label
        }
        records.append(rec)

    return pd.DataFrame(records)

def build_12month_matrix(df_wide, metric_col, days_dict):
    """
    Construct official PTIT 12-month matrix format with monthly columns,
    daily average, and cumulative total for a given metric.
    """
    active_months = [m for m in THAI_MONTHS if m in df_wide['เดือน'].unique()]
    
    # Pivot: Index = [PTIT_Region, PTIT_Operator_Field, PTIT_Order], Columns = เดือน
    pivot_df = df_wide.pivot_table(
        index=['PTIT_Region', 'PTIT_Operator_Field', 'PTIT_Order'],
        columns='เดือน',
        values=metric_col,
        aggfunc='sum'
    ).reset_index()

    # Re-order columns
    col_order = ['PTIT_Region', 'PTIT_Operator_Field', 'PTIT_Order'] + [m for m in THAI_MONTHS if m in pivot_df.columns]
    pivot_df = pivot_df[[c for c in col_order if c in pivot_df.columns]]
    
    # Calculate Year Average and Cumulative Total
    month_cols = [m for m in THAI_MONTHS if m in pivot_df.columns]
    total_active_days = sum(days_dict.get(m, 30) for m in month_cols)
    
    def calc_row_stats(row):
        total_vol = 0.0
        for m in month_cols:
            val = row.get(m, 0.0)
            if pd.notna(val):
                days = days_dict.get(m, 30)
                total_vol += float(val) * days
        avg_rate = total_vol / total_active_days if total_active_days > 0 else 0.0
        # For gas: BCF (Billion cubic feet) / For liquids: MMbbl (Million barrels)
        cum_metric = (total_vol / 1000.0) if 'ล้าน' in metric_col else (total_vol / 1_000_000.0)
        return pd.Series({'เฉลี่ยทั้งปี': avg_rate, 'สะสมทั้งปี': cum_metric})

    stats_df = pivot_df.apply(calc_row_stats, axis=1)
    pivot_df['เฉลี่ยทั้งปี'] = stats_df['เฉลี่ยทั้งปี']
    pivot_df['สะสมทั้งปี'] = stats_df['สะสมทั้งปี']

    # Sort by PTIT_Region (Onshore first), then PTIT_Order
    pivot_df['is_onshore'] = pivot_df['PTIT_Region'].apply(lambda x: 0 if str(x).strip().lower() == 'onshore' else 1)
    pivot_df = pivot_df.sort_values(by=['is_onshore', 'PTIT_Order']).drop(columns=['is_onshore'])
    pivot_df = pivot_df.fillna(0.0)
    
    return pivot_df

def run_etl():
    print(f"[START] PTIT Hydrocarbon ETL Pipeline...")
    t0 = time.time()
    
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    FRONTEND_DATA_DIR.mkdir(parents=True, exist_ok=True)
    EXCEL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Fetch Inventory
    print("[DMF] Discovering DMF online monthly inventories...")
    inv = dmf.get_dmf_online_inventory(timeout=10) if dmf else {'production': [], 'sales': [], 'is_live': False}
    print(f"   DMF Status: {'ONLINE (Live Portal)' if inv.get('is_live') else 'LOCAL REPO FALLBACK'}")
    print(f"   Production months available: {len(inv.get('production', []))}")
    print(f"   Sales months available: {len(inv.get('sales', []))}")

    # 2. Extract & Parse Production
    prod_dfs = []
    prod_items = sorted(inv.get('production', []), key=lambda x: (x.get('year_be', 2569), x.get('month', 1)))
    
    # Fallback to local files if no items in inventory
    if not prod_items:
        local_prod_files = sorted(glob.glob(str(REPO_ROOT / "createpetroleum_2569_*.xlsx")))
        for pf in local_prod_files:
            m = re.search(r'createpetroleum_(\d{4})_(\d+)\.xlsx', os.path.basename(pf))
            if m:
                prod_items.append({'year_be': int(m.group(1)), 'month': int(m.group(2)), 'label': os.path.basename(pf)})

    print(f"[PROD] Processing {len(prod_items)} production monthly records...")
    for it in prod_items:
        y_be = it.get('year_be', 2569)
        m_num = it.get('month', 1)
        lbl = it.get('label', f"{m_num}/{y_be}")
        try:
            if dmf:
                buf, fname = dmf.stream_dmf_production_bytes(y_be, m_num)
            else:
                loc = REPO_ROOT / f"createpetroleum_{y_be}_{m_num}.xlsx"
                with open(loc, 'rb') as f:
                    buf = io.BytesIO(f.read())
                fname = loc.name
            df_p = parse_excel_file(buf, f"DMF_{fname}")
            prod_dfs.append(df_p)
            print(f"   [OK] Parsed Production {lbl}: {len(df_p)} field rows")
        except Exception as e:
            print(f"   [WARN] Error loading Production {lbl}: {e}")

    # Process Production Flat Tables & Matrices
    prod_result = {}
    if prod_dfs:
        df_raw_prod = pd.concat(prod_dfs, ignore_index=True)
        df_master = load_master_mapping()
        
        master_cols = ['Lookup_Key', 'PTIT_Region', 'PTIT_Operator_Field', 'PTIT_Order', 'ผู้ดำเนินการ', 'แอ่งปิโตรเลียม', 'ประเภทสัญญา']
        master_cols = [c for c in master_cols if c in df_master.columns]
        
        df_merged_prod = pd.merge(df_raw_prod, df_master[master_cols], on='Lookup_Key', how='left')
        df_merged_prod['PTIT_Region'] = df_merged_prod['PTIT_Region'].fillna('Offshore')
        df_merged_prod['PTIT_Operator_Field'] = df_merged_prod['PTIT_Operator_Field'].fillna(df_merged_prod['แหล่ง_ไฟล์ดิบ'])
        df_merged_prod['ผู้ดำเนินการ'] = df_merged_prod['ผู้ดำเนินการ'].fillna('ยังไม่ระบุ')
        df_merged_prod['แอ่งปิโตรเลียม'] = df_merged_prod['แอ่งปิโตรเลียม'].fillna('ยังไม่ระบุ')
        df_merged_prod['ประเภทสัญญา'] = df_merged_prod['ประเภทสัญญา'].fillna('ยังไม่ระบุ')
        df_merged_prod['PTIT_Order'] = df_merged_prod['PTIT_Order'].fillna(99).astype(int)

        # Inject DEDP Fang Crude
        df_wide = inject_fang_to_dataframe(df_merged_prod)
        
        # Sort chronologically
        df_wide = df_wide.sort_values(
            by=['ปี', 'ลำดับเดือน', 'พื้นที่', 'PTIT_Order'],
            ascending=[True, True, False, True]
        ).reset_index(drop=True)

        active_months = [m for m in THAI_MONTHS if m in df_wide['เดือน'].unique()]
        days_dict = {m: get_days_in_month(m, df_wide['ปี'].iloc[0] if len(df_wide) > 0 else 2569) for m in active_months}
        
        # Build 12-Month Matrices
        metrics_map = {
            'gas_mmscfd': ('ก๊าซธรรมชาติ (ล้านลบ.ฟุต/วัน)', 'MMSCFD'),
            'condensate_bpd': ('ก๊าซธรรมชาติเหลว (บาร์เรล/วัน)', 'BPD'),
            'crude_bpd': ('น้ำมันดิบ (บาร์เรล/วัน)', 'BPD'),
            'total_boed': ('รวมเทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)', 'BOED')
        }
        matrices = {}
        for m_key, (col_name, unit_lbl) in metrics_map.items():
            if col_name in df_wide.columns:
                mat_df = build_12month_matrix(df_wide, col_name, days_dict)
                matrices[m_key] = {
                    'column_name': col_name,
                    'unit': unit_lbl,
                    'rows': mat_df.to_dict(orient='records')
                }

        # Calculate high-level Executive KPIs (Latest Month & Year Total)
        latest_month = df_wide.sort_values('ลำดับเดือน').iloc[-1]['เดือน'] if len(df_wide) > 0 else 'สิงหาคม'
        latest_df = df_wide[df_wide['เดือน'] == latest_month]

        kpis = {
            'latest_month': latest_month,
            'reporting_year_be': int(df_wide['ปี'].iloc[0]) if len(df_wide) > 0 else 2569,
            'reporting_year_ce': (int(df_wide['ปี'].iloc[0]) - 543) if len(df_wide) > 0 else 2026,
            'active_months': active_months,
            'latest_gas_mmscfd': round(float(latest_df['ก๊าซธรรมชาติ (ล้านลบ.ฟุต/วัน)'].sum()), 2),
            'latest_cond_bpd': round(float(latest_df['ก๊าซธรรมชาติเหลว (บาร์เรล/วัน)'].sum()), 2),
            'latest_crude_bpd': round(float(latest_df['น้ำมันดิบ (บาร์เรล/วัน)'].sum()), 2),
            'latest_boed': round(float(latest_df['รวมเทียบเท่าน้ำมันดิบ (บาร์เรล/วัน)'].sum()), 2),
            'total_active_fields': int(df_wide['PTIT_Operator_Field'].nunique()),
            'total_concessions': int(df_wide['แปลง_ไฟล์ดิบ'].nunique()),
            'total_operators': int(df_wide['ผู้ดำเนินการ'].nunique())
        }

        # Clean DataFrame to ensure strict standard JSON compliance (no NaN or inf)
        df_wide_json = df_wide.copy()
        for col in df_wide_json.select_dtypes(include=['float', 'int']).columns:
            df_wide_json[col] = df_wide_json[col].fillna(0.0)
        df_wide_json = df_wide_json.fillna('')

        prod_result = {
            'metadata': {
                'title': 'PTIT Focus Statistics - Domestic Petroleum Production',
                'source': 'Department of Mineral Fuels (DMF), Defence Energy Department (DEDP)',
                'publisher': 'Petroleum Institute of Thailand',
                'sync_timestamp': datetime.datetime.now().isoformat(),
                'record_count': len(df_wide),
                'active_months': active_months
            },
            'kpis': kpis,
            'matrices': matrices,
            'flat_records': df_wide_json.to_dict(orient='records')
        }
        prod_result = clean_for_json(prod_result)

        # Save to Excel Flat Table
        try:
            with pd.ExcelWriter(EXCEL_PROD_OUTPUT, engine='openpyxl') as writer:
                df_wide.to_excel(writer, sheet_name='Flat_Wide', index=False)
                df_master.to_excel(writer, sheet_name='Master_Mapping', index=False)
            print(f"   [OK] Wrote Excel: {EXCEL_PROD_OUTPUT}")
        except Exception as ex:
            print(f"   [WARN] Could not write Excel Flat: {ex}")

    # 3. Extract & Parse Sales
    sale_dfs = []
    sale_items = sorted(inv.get('sales', []), key=lambda x: (x.get('year_ce', 2026), x.get('month', 1)))
    
    if not sale_items:
        local_sale_files = sorted(glob.glob(str(REPO_ROOT / "Sale" / "salevalue_2026_*.xlsx")))
        for sf in local_sale_files:
            m = re.search(r'salevalue_(\d{4})_(\d+)\.xlsx', os.path.basename(sf))
            if m:
                sale_items.append({'year_ce': int(m.group(1)), 'month': int(m.group(2)), 'label': os.path.basename(sf)})

    print(f"[SALE] Processing {len(sale_items)} sales monthly records...")
    for it in sale_items:
        y_ce = it.get('year_ce', 2026)
        m_num = it.get('month', 1)
        lbl = it.get('label', f"{m_num}/{y_ce}")
        try:
            if dmf:
                buf, fname = dmf.stream_dmf_sales_bytes(y_ce, m_num)
            else:
                s_cand = REPO_ROOT / "Sale" / f"salevalue_{y_ce}_{m_num:02d}.xlsx"
                with open(s_cand, 'rb') as f:
                    buf = io.BytesIO(f.read())
                fname = s_cand.name
            df_s = parse_sales_file(buf, f"DMF_{fname}")
            sale_dfs.append(df_s)
            print(f"   [OK] Parsed Sales {lbl}: {len(df_s)} product rows")
        except Exception as e:
            print(f"   [WARN] Error loading Sales {lbl}: {e}")

    sales_result = {}
    if sale_dfs:
        df_raw_sales = pd.concat(sale_dfs, ignore_index=True)
        df_sale_master = load_sale_master_mapping()
        
        df_sales_merged = pd.merge(df_raw_sales, df_sale_master, on='แหล่ง_ไฟล์ดิบ', how='left')
        df_sales_merged['พื้นที่'] = df_sales_merged['พื้นที่'].fillna('ไม่ระบุ')
        df_sales_merged['ผู้ดำเนินการ'] = df_sales_merged['ผู้ดำเนินการ'].fillna('ไม่ระบุ')
        df_sales_merged['แอ่งปิโตรเลียม'] = df_sales_merged['แอ่งปิโตรเลียม'].fillna('ไม่ระบุ')
        df_sales_merged['ประเภทสัญญา'] = df_sales_merged['ประเภทสัญญา'].fillna('ไม่ระบุ')

        # Compute price wellhead
        if 'ราคาปากหลุม_Wellhead_Price' not in df_sales_merged.columns:
            df_sales_merged['ราคาปากหลุม_Wellhead_Price'] = 0.0
            mask_gas = (df_sales_merged['ประเภทปิโตรเลียม'] == 'ก๊าซธรรมชาติ') & (df_sales_merged.get('ปริมาณการขาย_MMBTU', 0) > 0)
            df_sales_merged.loc[mask_gas, 'ราคาปากหลุม_Wellhead_Price'] = df_sales_merged.loc[mask_gas, 'มูลค่าการขาย_บาท'] / df_sales_merged.loc[mask_gas, 'ปริมาณการขาย_MMBTU']
            mask_oil = (df_sales_merged['ประเภทปิโตรเลียม'].isin(['ก๊าซธรรมชาติเหลว', 'น้ำมันดิบ'])) & (df_sales_merged.get('ปริมาณการขาย_บาร์เรล', 0) > 0)
            df_sales_merged.loc[mask_oil, 'ราคาปากหลุม_Wellhead_Price'] = df_sales_merged.loc[mask_oil, 'มูลค่าการขาย_บาท'] / df_sales_merged.loc[mask_oil, 'ปริมาณการขาย_บาร์เรล']

        df_sales_flat = df_sales_merged.sort_values(['ลำดับเดือน', 'ประเภทปิโตรเลียม', 'แหล่ง_ไฟล์ดิบ']).reset_index(drop=True)
        
        total_val_thb = float(df_sales_flat['มูลค่าการขาย_บาท'].sum())
        total_royalty_thb = float(df_sales_flat['ค่าภาคหลวง_บาท'].sum())
        avg_royalty_pct = (total_royalty_thb / total_val_thb * 100) if total_val_thb > 0 else 0.0

        # Clean DataFrame to ensure strict standard JSON compliance (no NaN or inf)
        df_sales_json = df_sales_flat.copy()
        for col in df_sales_json.select_dtypes(include=['float', 'int']).columns:
            df_sales_json[col] = df_sales_json[col].fillna(0.0)
        df_sales_json = df_sales_json.fillna('')

        sales_result = {
            'metadata': {
                'title': 'PTIT Focus Statistics - Domestic Petroleum Sales & Royalty',
                'source': 'Department of Mineral Fuels (DMF)',
                'publisher': 'Petroleum Institute of Thailand',
                'sync_timestamp': datetime.datetime.now().isoformat(),
                'record_count': len(df_sales_flat)
            },
            'kpis': {
                'total_sales_value_thb': total_val_thb,
                'total_sales_value_million_thb': round(total_val_thb / 1_000_000.0, 2),
                'total_royalty_thb': total_royalty_thb,
                'total_royalty_million_thb': round(total_royalty_thb / 1_000_000.0, 2),
                'effective_royalty_rate_pct': round(avg_royalty_pct, 2)
            },
            'flat_records': df_sales_json.to_dict(orient='records')
        }
        sales_result = clean_for_json(sales_result)

        # Save to Excel
        try:
            with pd.ExcelWriter(EXCEL_SALE_OUTPUT, engine='openpyxl') as writer:
                df_sales_flat.to_excel(writer, sheet_name='Sale_Flat_Table', index=False)
                df_sale_master.to_excel(writer, sheet_name='Sale_Master_Mapping', index=False)
            print(f"   [OK] Wrote Excel: {EXCEL_SALE_OUTPUT}")
        except Exception as ex:
            print(f"   [WARN] Could not write Sales Excel Flat: {ex}")

    # 4. Save JSON files
    for target_dir in [DATA_DIR, FRONTEND_DATA_DIR]:
        target_dir.mkdir(parents=True, exist_ok=True)
        with open(target_dir / "production_master.json", "w", encoding="utf-8") as f:
            json.dump(prod_result, f, ensure_ascii=False, indent=2)
        with open(target_dir / "sales_master.json", "w", encoding="utf-8") as f:
            json.dump(sales_result, f, ensure_ascii=False, indent=2)

    manifest = {
        'timestamp': datetime.datetime.now().isoformat(),
        'execution_duration_sec': round(time.time() - t0, 2),
        'status': 'success',
        'is_live_dmf': inv.get('is_live', False),
        'production_records': len(prod_result.get('flat_records', [])),
        'sales_records': len(sales_result.get('flat_records', [])),
        'latest_production_month': prod_result.get('kpis', {}).get('latest_month', 'N/A'),
        'checksum_prod': hashlib.md5(json.dumps(prod_result).encode('utf-8')).hexdigest(),
        'checksum_sale': hashlib.md5(json.dumps(sales_result).encode('utf-8')).hexdigest()
    }

    for target_dir in [DATA_DIR, FRONTEND_DATA_DIR]:
        with open(target_dir / "sync_manifest.json", "w", encoding="utf-8") as f:
            json.dump(manifest, f, ensure_ascii=False, indent=2)

    print(f"[SUCCESS] ETL Pipeline completed in {manifest['execution_duration_sec']}s!")
    print(f"   Production Records: {manifest['production_records']}")
    print(f"   Sales Records: {manifest['sales_records']}")
    print(f"   Manifest: {DATA_DIR / 'sync_manifest.json'}")

if __name__ == '__main__':
    run_etl()
